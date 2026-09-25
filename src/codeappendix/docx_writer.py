"""Build the Word document.

Headings, the summary table and page setup go through python-docx. Code lines
are the bulk of the document, so they are emitted as raw WordprocessingML,
which is an order of magnitude faster than building them run by run.
"""

from __future__ import annotations

import hashlib
import math
from datetime import datetime, timezone
from pathlib import Path
from typing import List, Optional, Sequence
from xml.sax.saxutils import escape

from docx import Document
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls, qn
from docx.shared import Cm, Pt, RGBColor

from . import __version__, layout
from .config import Config
from .highlight import Highlighter, Segment
from .i18n import fmt_int, t
from .reader import SourceFile

LINE_NUMBER_COLOR = RGBColor(0x80, 0x80, 0x80)
MUTED_COLOR = RGBColor(0x59, 0x59, 0x59)
HEADER_FILL = "E7E6E6"


def _set_fonts(rPr, name: str) -> None:
    rFonts = rPr.get_or_add_rFonts()
    for attr in ("asciiTheme", "hAnsiTheme", "eastAsiaTheme", "cstheme"):
        key = qn(f"w:{attr}")
        if key in rFonts.attrib:
            del rFonts.attrib[key]
    for attr in ("ascii", "hAnsi", "eastAsia", "cs"):
        rFonts.set(qn(f"w:{attr}"), name)


def _set_color(style, rgb: RGBColor) -> None:
    style.font.color.rgb = rgb
    color = style.element.rPr.find(qn("w:color"))
    for attr in ("themeColor", "themeShade", "themeTint"):
        key = qn(f"w:{attr}")
        if key in color.attrib:
            del color.attrib[key]


def _para_style(
    doc,
    name: str,
    *,
    base: Optional[str] = None,
    font: str,
    size: float,
    bold: Optional[bool] = None,
    italic: Optional[bool] = None,
    color: Optional[RGBColor] = None,
    align=None,
    before: float = 0,
    after: float = 0,
    line_spacing: float = 1.0,
    keep_next: bool = False,
    indent_pt: float = 0,
    no_proof: bool = False,
):
    styles = doc.styles
    try:
        style = styles[name]
    except KeyError:
        style = styles.add_style(name, WD_STYLE_TYPE.PARAGRAPH)
        style.base_style = styles[base] if base else None
        style.quick_style = False
    rPr = style.element.get_or_add_rPr()
    _set_fonts(rPr, font)
    style.font.size = Pt(size)
    if bold is not None:
        style.font.bold = bold
    if italic is not None:
        style.font.italic = italic
    if color is not None:
        _set_color(style, color)
    if no_proof:
        rPr.get_or_add_noProof()
    pf = style.paragraph_format
    pf.space_before = Pt(before)
    pf.space_after = Pt(after)
    pf.line_spacing = line_spacing
    pf.keep_with_next = keep_next
    if align is not None:
        pf.alignment = align
    if indent_pt:
        pf.left_indent = Pt(indent_pt)
        pf.first_line_indent = Pt(-indent_pt)
    return style


class DocxBuilder:
    def __init__(self, cfg: Config, project_name: str) -> None:
        self.cfg = cfg
        self.lang = cfg.language
        self.project_name = project_name
        self.doc = Document()
        self.highlighter = Highlighter(cfg.highlight, cfg.theme)
        self._code_styles: dict = {}
        self._bookmark_prefix = bookmark_prefix(cfg, project_name)
        self._setup_page()
        self._setup_compat()
        self._setup_styles()

    # ----------------------------------------------------------------- setup
    def _setup_page(self) -> None:
        cfg = self.cfg
        width, height = layout.PAGE_SIZES_CM[cfg.page_size]
        section = self.doc.sections[0]
        section.page_width = Cm(width)
        section.page_height = Cm(height)
        section.top_margin = Cm(cfg.margin_top)
        section.bottom_margin = Cm(cfg.margin_bottom)
        section.left_margin = Cm(cfg.margin_left)
        section.right_margin = Cm(cfg.margin_right)
        section.header_distance = Cm(min(1.25, cfg.margin_top / 2))
        section.footer_distance = Cm(min(1.25, cfg.margin_bottom / 2))
        if cfg.page_numbers:
            p = section.footer.paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p._p.append(parse_xml(
                f'<w:fldSimple {nsdecls("w")} w:instr=" PAGE "><w:r><w:t>1</w:t></w:r></w:fldSimple>'
            ))

        props = self.doc.core_properties
        props.title = cfg.effective_title
        props.subject = self.project_name
        props.author = ""
        props.comments = f"CodeAppendix {__version__}"
        # python-docx writes these as UTC ("...Z"), so give it UTC.
        props.created = props.modified = datetime.now(timezone.utc).replace(tzinfo=None)

    def _setup_compat(self) -> None:
        """Mark the file as a Word 2013+ document (no "Compatibility Mode")."""
        for setting in self.doc.settings.element.iter(qn("w:compatSetting")):
            if setting.get(qn("w:name")) == "compatibilityMode":
                setting.set(qn("w:val"), "15")

    def _setup_styles(self) -> None:
        cfg, doc = self.cfg, self.doc
        # Document defaults: explicit fonts (no theme fonts), language for the
        # spell checker, and no inherited spacing.
        defaults = doc.styles.element.find(qn("w:docDefaults"))
        rPr = defaults.find(qn("w:rPrDefault")).find(qn("w:rPr"))
        _set_fonts(rPr, cfg.font)
        lang = rPr.find(qn("w:lang"))
        lang.set(qn("w:val"), "vi-VN" if self.lang == "vi" else "en-US")
        spacing = defaults.find(qn("w:pPrDefault")).find(qn("w:pPr")).find(qn("w:spacing"))
        spacing.set(qn("w:after"), "0")
        spacing.set(qn("w:line"), "240")

        _para_style(doc, "Normal", font=cfg.font, size=cfg.font_size,
                    after=6, line_spacing=cfg.line_spacing)
        _para_style(doc, "Heading 1", font=cfg.font, size=cfg.font_size + 1,
                    bold=True, color=RGBColor(0, 0, 0), align=WD_ALIGN_PARAGRAPH.CENTER,
                    before=0, after=12, line_spacing=cfg.line_spacing, keep_next=True)
        _para_style(doc, "Heading 2", font=cfg.font, size=cfg.font_size,
                    bold=True, color=RGBColor(0, 0, 0), before=12, after=4,
                    line_spacing=1.0, keep_next=True)
        # Same look as Heading 2 but outside the outline, so it stays out of
        # the thesis table of contents.
        _para_style(doc, "Code File", base="Normal", font=cfg.font, size=cfg.font_size,
                    bold=True, before=12, after=4, line_spacing=1.0, keep_next=True)
        _para_style(doc, "Code Caption", base="Normal", font=cfg.font, size=cfg.font_size,
                    bold=True, before=12, after=6, line_spacing=1.0, keep_next=True)
        _para_style(doc, "Code File Info", base="Normal", font=cfg.font,
                    size=max(8, cfg.font_size - 2), italic=True, color=MUTED_COLOR,
                    after=4, line_spacing=1.0, keep_next=True)
        _para_style(doc, "Code Tree", font=cfg.code_font, size=cfg.code_size,
                    line_spacing=1.0, no_proof=True)
        _para_style(doc, "Code Omitted", font=cfg.font, size=max(8, cfg.font_size - 2),
                    italic=True, color=MUTED_COLOR, align=WD_ALIGN_PARAGRAPH.CENTER,
                    before=2, after=2, line_spacing=1.0)
        _para_style(doc, "Code Table", font=cfg.font, size=max(8, cfg.font_size - 1),
                    before=1, after=1, line_spacing=1.0)

        num = doc.styles.add_style("Code Line Number", WD_STYLE_TYPE.CHARACTER)
        num.font.color.rgb = LINE_NUMBER_COLOR
        num.quick_style = False

    def _code_style(self, digits: int) -> str:
        """Paragraph style for code whose line numbers have ``digits`` digits."""
        key = digits if self.cfg.line_numbers else 0
        if key not in self._code_styles:
            style = _para_style(
                self.doc,
                f"Code {key}" if key else "Code",
                font=self.cfg.code_font,
                size=self.cfg.code_size,
                line_spacing=1.0,
                indent_pt=layout.code_indent_pt(self.cfg, key),
                no_proof=True,
            )
            self._code_styles[key] = style.style_id
        return self._code_styles[key]

    # --------------------------------------------------------------- content
    def heading_label(self, index: int) -> str:
        return f"{self.cfg.label}.{index}." if self.cfg.label else f"{index}."

    def build(self, files: Sequence[SourceFile], tree_lines: Optional[List[str]]) -> None:
        cfg, lang = self.cfg, self.lang
        body = self.doc
        body.add_paragraph(cfg.effective_title, style="Heading 1")

        total_lines = sum(f.total_lines for f in files)
        intro = t(lang, "intro", project=self.project_name, files=fmt_int(lang, len(files)),
                  lines=fmt_int(lang, total_lines))
        if any(f.omitted for f in files):
            intro += t(lang, "intro_truncated", max=fmt_int(lang, cfg.max_lines))
        p = body.add_paragraph(intro)
        p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY

        if cfg.summary_table:
            self._summary_table(files, total_lines)
        if tree_lines:
            body.add_paragraph(t(lang, "tree_caption"), style="Code Caption")
            style_id = self.doc.styles["Code Tree"].style_id
            # Short trees stay on one page; long ones may split.
            keep_together = len(tree_lines) < 60
            self._append_xml("".join(
                f'<w:p><w:pPr><w:pStyle w:val="{style_id}"/>'
                f'{"<w:keepNext/>" if keep_together and i < len(tree_lines) - 1 else ""}'
                f"</w:pPr>{_run(line)}</w:p>"
                for i, line in enumerate(tree_lines)
            ))

        for index, src in enumerate(files, start=1):
            self._file(index, src)

        if cfg.summary_table:
            self._request_field_update()

    def _summary_table(self, files: Sequence[SourceFile], total_lines: int) -> None:
        lang = self.lang
        self.doc.add_paragraph(t(lang, "list_caption"), style="Code Caption")
        headers = [t(lang, k) for k in ("col_no", "col_file", "col_language", "col_lines", "col_page")]
        table = self.doc.add_table(rows=1, cols=len(headers))
        table.style = self.doc.styles["Table Grid"]
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        table.autofit = False

        usable_cm = layout.usable_width_pt(self.cfg) / layout.CM_TO_PT
        widths = [1.2, 0.0, 3.4, 1.9, 1.6]
        widths[1] = max(4.0, usable_cm - sum(widths))
        align = {
            0: WD_ALIGN_PARAGRAPH.CENTER,
            3: WD_ALIGN_PARAGRAPH.RIGHT,
            4: WD_ALIGN_PARAGRAPH.CENTER,
        }
        cell_style = self.doc.styles["Code Table"]

        def fill(cells, values, bold: bool = False) -> None:
            for i, (cell, value) in enumerate(zip(cells, values)):
                cell.width = Cm(widths[i])
                para = cell.paragraphs[0]
                para.style = cell_style
                para.alignment = WD_ALIGN_PARAGRAPH.CENTER if bold else align.get(i)
                if isinstance(value, str):
                    para.add_run(value).bold = bold
                else:
                    para._p.append(value)

        # Word reads column widths from the cells, LibreOffice from the grid;
        # set both (``autofit = False`` already fixed the layout).
        for grid_col, width in zip(table._tbl.tblGrid.findall(qn("w:gridCol")), widths):
            grid_col.set(qn("w:w"), str(int(Cm(width).twips)))

        header = table.rows[0]
        fill(header.cells, headers, bold=True)
        for cell in header.cells:
            cell._tc.get_or_add_tcPr().append(parse_xml(
                f'<w:shd {nsdecls("w")} w:val="clear" w:color="auto" w:fill="{HEADER_FILL}"/>'
            ))
        # Repeat the header row on every page the table spans.
        header._tr.get_or_add_trPr().append(parse_xml(f'<w:tblHeader {nsdecls("w")}/>'))

        for index, src in enumerate(files, start=1):
            row = table.add_row()
            page_ref = parse_xml(
                f'<w:fldSimple {nsdecls("w")} w:instr=" PAGEREF {self._bookmark(index)} \\h ">'
                f"<w:r><w:t></w:t></w:r></w:fldSimple>"
            )
            fill(row.cells, [str(index), breakable(src.rel), src.language,
                             fmt_int(lang, src.total_lines), page_ref])
            row._tr.get_or_add_trPr().append(parse_xml(f'<w:cantSplit {nsdecls("w")}/>'))

        total = table.add_row()
        fill(total.cells, ["", "", "", fmt_int(lang, total_lines), ""], bold=True)
        total.cells[3].paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.RIGHT
        label = total.cells[0].merge(total.cells[2])
        label.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.LEFT
        label.paragraphs[0].add_run(t(lang, "total")).bold = True
        for extra in label.paragraphs[1:]:
            extra._p.getparent().remove(extra._p)

    def _file(self, index: int, src: SourceFile) -> None:
        cfg, lang = self.cfg, self.lang
        styles = self.doc.styles
        heading_style = styles["Heading 2" if cfg.file_heading == "heading" else "Code File"]
        name = self._bookmark(index)
        page_break = '<w:pageBreakBefore/>' if cfg.page_break_per_file and index > 1 else ""
        xml = [
            f'<w:p><w:pPr><w:pStyle w:val="{heading_style.style_id}"/>{page_break}</w:pPr>'
            f'<w:bookmarkStart w:id="{index}" w:name="{name}"/>'
            f"{_run(f'{self.heading_label(index)} {breakable(src.rel)}')}"
            f'<w:bookmarkEnd w:id="{index}"/></w:p>'
        ]

        head_count = src.omitted[0] - 1 if src.omitted else len(src.lines)
        head_rows, _ = self.highlighter.lines(src.lines[:head_count], src.lexer, src.rel)
        tail_rows: List[List[Segment]] = []
        if src.omitted:
            tail_rows, _ = self.highlighter.lines(src.lines[head_count:], src.lexer, src.rel)

        if cfg.file_info:
            key = "file_info_truncated" if src.omitted else "file_info"
            info = t(lang, key, language=src.language, lines=fmt_int(lang, src.total_lines),
                     shown=fmt_int(lang, len(src.lines)))
            xml.append(f'<w:p><w:pPr><w:pStyle w:val="{styles["Code File Info"].style_id}"/></w:pPr>{_run(info)}</w:p>')

        digits = layout.number_width(src.total_lines)
        code_style = self._code_style(digits)
        xml.extend(self._code_paragraphs(head_rows, 1, digits, code_style))
        if src.omitted:
            start, count = src.omitted
            note = t(lang, "omitted", count=fmt_int(lang, count), start=fmt_int(lang, start),
                     end=fmt_int(lang, start + count - 1))
            xml.append(f'<w:p><w:pPr><w:pStyle w:val="{styles["Code Omitted"].style_id}"/></w:pPr>{_run(note)}</w:p>')
            xml.extend(self._code_paragraphs(tail_rows, start + count, digits, code_style))
        self._append_xml("".join(xml))

    def _code_paragraphs(self, rows: List[List[Segment]], first: int, digits: int, style_id: str):
        numbers = self.cfg.line_numbers
        num_style = self.doc.styles["Code Line Number"].style_id
        # Keep the first two lines with the heading, and never leave a single
        # line stranded at the end of a file.
        last = len(rows) - 1
        for i, segments in enumerate(rows):
            keep = "<w:keepNext/>" if i < 1 or (last >= 2 and i == last - 1) else ""
            parts = [f'<w:p><w:pPr><w:pStyle w:val="{style_id}"/>{keep}</w:pPr>']
            if numbers:
                parts.append(
                    f'<w:r><w:rPr><w:rStyle w:val="{num_style}"/></w:rPr>'
                    f'<w:t xml:space="preserve">{str(first + i).rjust(digits)}</w:t><w:tab/></w:r>'
                )
            for text, (color, bold, italic) in segments:
                props = ("<w:b/>" if bold else "") + ("<w:i/>" if italic else "") + (
                    f'<w:color w:val="{color}"/>' if color else ""
                )
                parts.append(_run(text, props))
            parts.append("</w:p>")
            yield "".join(parts)

    # --------------------------------------------------------------- helpers
    def _bookmark(self, index: int) -> str:
        return f"{self._bookmark_prefix}{index}"

    def _append_xml(self, paragraphs_xml: str) -> None:
        if not paragraphs_xml:
            return
        container = parse_xml(f"<w:body {nsdecls('w')}>{paragraphs_xml}</w:body>")
        body = self.doc.element.body
        sect = body.sectPr
        for child in list(container):
            if sect is not None:
                sect.addprevious(child)
            else:
                body.append(child)

    def _request_field_update(self) -> None:
        """Ask Word to refresh fields (the page column) when the file opens."""
        settings = self.doc.settings.element
        if settings.find(qn("w:updateFields")) is not None:
            return
        element = parse_xml(f'<w:updateFields {nsdecls("w")} w:val="true"/>')
        compat = settings.find(qn("w:compat"))
        if compat is not None:
            compat.addprevious(element)
        else:
            settings.append(element)

    def save(self, path: Path) -> None:
        self.doc.save(str(path))


def breakable(path: str) -> str:
    """Let long paths wrap after each "/" instead of mid-word."""
    return path.replace("/", "/\u200b")


def bookmark_prefix(cfg: Config, project_name: str) -> str:
    """Bookmark prefix unique to this appendix.

    Two appendices pasted into one thesis must not share bookmark names, or
    the page column of one would point into the other.
    """
    key = f"{project_name}|{cfg.effective_title}|{cfg.label}".encode("utf-8")
    # A leading underscore makes Word treat the bookmark as hidden.
    return f"_CA{hashlib.sha1(key).hexdigest()[:8]}_"


def _run(text: str, props: str = "") -> str:
    rpr = f"<w:rPr>{props}</w:rPr>" if props else ""
    return f'<w:r>{rpr}<w:t xml:space="preserve">{escape(text)}</w:t></w:r>'


def estimate_pages(cfg: Config, files: Sequence[SourceFile], tree_lines: Optional[List[str]]) -> int:
    """Rough page count, good to about ±10 %."""
    usable_h = layout.usable_height_pt(cfg)
    code_h = layout.code_line_height_pt(cfg)
    text_h = layout.text_line_height_pt(cfg, cfg.font_size)
    small_h = layout.text_line_height_pt(cfg, max(8, cfg.font_size - 1))

    height = text_h * 1.2 + 12 + text_h * 3 + 6  # title + intro
    if cfg.summary_table:
        height += 18 + small_h * 1.1 * (len(files) + 2)
    if tree_lines:
        height += 18 + code_h * len(tree_lines)

    pages = 0.0
    for src in files:
        digits = layout.number_width(src.total_lines)
        block = 16 + text_h  # heading
        if cfg.file_info:
            block += small_h + 4
        block += code_h * layout.wrapped_rows(src.lines, cfg, digits)
        if src.omitted:
            block += small_h + 4
        if cfg.page_break_per_file:
            pages += math.ceil(height / usable_h) if height else 0
            height = 0
        height += block
    pages += height / usable_h
    return max(1, math.ceil(pages))
