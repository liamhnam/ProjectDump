from __future__ import annotations

import os

import docx
import pytest

from codeappendix import __version__
from codeappendix.cli import main, slug
from codeappendix.config import CONFIG_FILENAME
from codeappendix.pdf import count_pages, find_soffice


@pytest.fixture
def project(make_project):
    return make_project({
        "app.py": "print('xin chào')\n",
        "web/index.ts": "export const x: number = 1;\n",
        "tests/test_app.py": "def test(): pass\n",
        "notes.md": "# ghi chú\n",
    }, name="Đồ án Tốt nghiệp")


def test_build_default_command_and_output_name(project, tmp_path, monkeypatch, capsys):
    monkeypatch.chdir(tmp_path)
    assert main([str(project)]) == 0
    out = tmp_path / "do-an-tot-nghiep-ma-nguon.docx"
    assert out.is_file()
    printed = capsys.readouterr().out
    assert "3 tệp" in printed and "Đã tạo" in printed
    assert "không phải mã nguồn 1" in printed


def test_build_with_options(project, tmp_path, capsys):
    out = tmp_path / "custom" / "phu-luc"
    code = main(["build", str(project), "-o", str(out), "-x", "tests/", "--lang", "en",
                 "--label", "B", "--highlight", "mono", "-q"])
    assert code == 0
    assert capsys.readouterr().out.strip() == ""
    doc = docx.Document(str(out.with_suffix(".docx")))
    assert doc.paragraphs[0].text == "APPENDIX B. SOURCE CODE"
    assert len(doc.tables[0].rows) == 2 + 2


def test_config_file_is_used_and_cli_lists_extend_it(project, tmp_path):
    (project / CONFIG_FILENAME).write_text(
        '[files]\nexclude = ["tests/"]\n[document]\nlabel = "C"\n', encoding="utf-8"
    )
    out = tmp_path / "a.docx"
    assert main([str(project), "-o", str(out), "-x", "web/", "-q"]) == 0
    doc = docx.Document(str(out))
    headings = [p.text.replace("​", "") for p in doc.paragraphs if p.style.name == "Heading 2"]
    assert headings == ["C.1. app.py"]


def test_list(project, capsys):
    assert main(["list", str(project)]) == 0
    out = capsys.readouterr().out
    assert "tests/test_app.py" in out and "web/index.ts" in out
    assert "không phải mã nguồn (1): notes.md" in out


def test_list_verbose(project, capsys):
    main(["list", str(project), "--verbose", "--lang", "en"])
    out = capsys.readouterr().out
    assert "not source code (1):" in out and "      notes.md" in out


def test_init(project, capsys):
    assert main(["init", str(project), "--preset", "compact"]) == 0
    text = (project / CONFIG_FILENAME).read_text(encoding="utf-8")
    assert 'preset = "compact"' in text
    assert main(["init", str(project)]) == 1
    assert "--force" in capsys.readouterr().err
    assert main(["init", str(project), "--force"]) == 0


def test_errors(tmp_path, capsys):
    assert main([str(tmp_path / "missing")]) == 1
    assert "Không tìm thấy thư mục" in capsys.readouterr().err
    empty = tmp_path / "empty"
    empty.mkdir()
    assert main([str(empty), "-o", str(tmp_path / "x.docx")]) == 1
    assert "Không tìm thấy tệp mã nguồn" in capsys.readouterr().err
    (empty / "a.py").write_text("x = 1\n")
    (empty / CONFIG_FILENAME).write_text("[code]\nhighlight = 'rainbow'\n")
    assert main([str(empty)]) == 1
    assert "highlight" in capsys.readouterr().err


def test_version(capsys):
    with pytest.raises(SystemExit):
        main(["--version"])
    assert __version__ in capsys.readouterr().out


def test_slug():
    assert slug("Đồ án Tốt nghiệp") == "do-an-tot-nghiep"
    assert slug("my_app 2") == "my-app-2"
    assert slug("***") == "project"


@pytest.mark.skipif(find_soffice() is None, reason="LibreOffice not installed")
def test_pdf_export(project, tmp_path, capsys):
    out = tmp_path / "a.docx"
    assert main([str(project), "-o", str(out), "--pdf"]) == 0
    pdf = out.with_suffix(".pdf")
    assert pdf.is_file() and os.path.getsize(pdf) > 1000
    assert count_pages(pdf) >= 1
    assert "PDF" in capsys.readouterr().out
