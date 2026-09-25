from __future__ import annotations

import re
import zipfile

import docx
import pytest

from codeappendix.config import Config, apply
from codeappendix.docx_writer import DocxBuilder, estimate_pages
from codeappendix.reader import load_all
from codeappendix.scanner import scan
from codeappendix.tree import render

ZWSP = "​"


def build(root, tmp_path, **settings):
    cfg = apply(Config(), {"project_name": root.name, **settings})
    result = scan(root, cfg)
    files = load_all(result, cfg)
    tree = render(root.name, [f.rel for f in files]) if cfg.tree != "none" else None
    builder = DocxBuilder(cfg, cfg.project_name)
    builder.build(files, tree)
    out = tmp_path / "out.docx"
    builder.save(out)
    return out, files


def xml(path, part="word/document.xml"):
    with zipfile.ZipFile(path) as z:
        return z.read(part).decode("utf-8")


def paragraphs(path, style_prefix):
    doc = docx.Document(str(path))
    return [p for p in doc.paragraphs if p.style.name.startswith(style_prefix)]


def code_lines(path):
    """Text of the code paragraphs (styles "Code 1", "Code 2", ...)."""
    doc = docx.Document(str(path))
    return [p.text for p in doc.paragraphs if re.fullmatch(r"Code \d+", p.style.name)]


@pytest.fixture
def project(make_project):
    return make_project({
        "src/app.py": 'def xin_chao(ten):\n    """Chào người dùng."""\n    return f"Xin chào {ten}"\n',
        "src/util.js": "// tiện ích\nexport const add = (a, b) => a + b;\n",
        "main.py": "from src.app import xin_chao\n\nprint(xin_chao('An'))\n",
    }, name="quan-ly-thu-vien")


def test_structure(project, tmp_path):
    out, files = build(project, tmp_path)
    doc = docx.Document(str(out))
    headings = [p.text.replace(ZWSP, "") for p in doc.paragraphs if p.style.name.startswith("Heading")]
    assert headings == [
        "PHỤ LỤC A. MÃ NGUỒN CHƯƠNG TRÌNH",
        "A.1. src/app.py",
        "A.2. src/util.js",
        "A.3. main.py",
    ]
    intro = doc.paragraphs[1].text
    assert "quan-ly-thu-vien" in intro and "3 tệp" in intro and "8 dòng" in intro

    lines = code_lines(out)
    assert len(lines) == sum(f.total_lines for f in files) == 8
    assert lines[0] == "1\tdef xin_chao(ten):"
    assert lines[1] == '2\t    """Chào người dùng."""'


def test_summary_table_with_page_references(project, tmp_path):
    out, _ = build(project, tmp_path)
    doc = docx.Document(str(out))
    table = doc.tables[0]
    rows = [[c.text.replace(ZWSP, "") for c in r.cells] for r in table.rows]
    assert rows[0] == ["STT", "Tệp", "Ngôn ngữ", "Số dòng", "Trang"]
    assert rows[1][:4] == ["1", "src/app.py", "Python", "3"]
    assert rows[2][:4] == ["2", "src/util.js", "JavaScript", "2"]
    assert rows[-1][0] == "Tổng cộng" and rows[-1][3] == "8"
    body = xml(out)
    names = re.findall(r'w:bookmarkStart w:id="\d+" w:name="(_CA[0-9a-f]{8}_\d+)"', body)
    assert len(names) == 3
    for name in names:
        assert f"PAGEREF {name} \\h" in body
    assert "<w:updateFields" in xml(out, "word/settings.xml")


def test_bookmarks_differ_between_appendices(project, tmp_path):
    a = xml(build(project, tmp_path, label="A")[0])
    b = xml(build(project, tmp_path, label="B")[0])
    name = re.compile(r'w:name="(_CA[0-9a-f]{8})_1"')
    assert name.search(a).group(1) != name.search(b).group(1)


def test_no_summary_no_field_update(project, tmp_path):
    out, _ = build(project, tmp_path, summary_table=False)
    assert not docx.Document(str(out)).tables
    assert "updateFields" not in xml(out, "word/settings.xml")


def test_styles_use_real_fonts_and_no_spellcheck_on_code(project, tmp_path):
    out, _ = build(project, tmp_path)
    styles = xml(out, "word/styles.xml")
    heading = re.search(r'w:styleId="Heading1".*?</w:style>', styles, re.S).group(0)
    assert "Theme" not in heading and 'w:ascii="Times New Roman"' in heading
    assert 'w:color w:val="000000"' in heading
    code = re.search(r'w:styleId="Code1".*?</w:style>', styles, re.S).group(0)
    assert "<w:noProof/>" in code and 'w:ascii="Consolas"' in code
    assert 'w:hanging="330"' in code  # (1 digit + 2) * 10pt * 0.55 = 16.5pt
    assert 'w:lang w:val="vi-VN"' in styles


def test_color_mono_none(project, tmp_path):
    common = {"summary_table": False, "file_info": False}
    colored = xml(build(project, tmp_path, highlight="color", **common)[0])
    assert '<w:color w:val="0000FF"/>' in colored  # vs theme keyword blue
    mono = xml(build(project, tmp_path, highlight="mono", **common)[0])
    assert "<w:color " not in mono and "<w:b/>" in mono and "<w:i/>" in mono
    plain = xml(build(project, tmp_path, highlight="none", **common)[0])
    assert "<w:color " not in plain and "<w:b/>" not in plain


def test_no_line_numbers(project, tmp_path):
    out, _ = build(project, tmp_path, line_numbers=False)
    code = [p for p in paragraphs(out, "Code") if p.style.name == "Code"]
    assert code[0].text == "def xin_chao(ten):"


def test_plain_headings_page_breaks_and_english(project, tmp_path):
    out, _ = build(project, tmp_path, file_heading="plain", page_break_per_file=True,
                   language="en", label="")
    doc = docx.Document(str(out))
    assert [p.text for p in doc.paragraphs if p.style.name == "Heading 2"] == []
    files = [p.text.replace(ZWSP, "") for p in doc.paragraphs if p.style.name == "Code File"]
    assert files == ["1. src/app.py", "2. src/util.js", "3. main.py"]
    assert doc.paragraphs[0].text == "SOURCE CODE"
    assert xml(out).count("<w:pageBreakBefore/>") == 2
    assert doc.tables[0].rows[0].cells[0].text == "No."


def test_truncation_numbers_tail_correctly(make_project, tmp_path):
    root = make_project({"long.py": "\n".join(f"v{i} = {i}" for i in range(1, 51)) + "\n"})
    out, _ = build(root, tmp_path, max_lines=6)
    # Numbers are right-aligned to the widest one (two digits here).
    assert code_lines(out) == [
        " 1\tv1 = 1", " 2\tv2 = 2", " 3\tv3 = 3", "48\tv48 = 48", "49\tv49 = 49", "50\tv50 = 50",
    ]
    omitted = [p.text for p in docx.Document(str(out)).paragraphs if p.style.name == "Code Omitted"]
    assert omitted == ["… lược bỏ 44 dòng (dòng 4–47) …"]


def test_special_characters_are_escaped(make_project, tmp_path):
    root = make_project({"a.html": "<p class=\"x\">A & B</p>\x01\n"})
    out, _ = build(root, tmp_path)
    # python-docx parses it, so the XML is well formed.
    assert code_lines(out) == ['1\t<p class="x">A & B</p>']


def test_page_setup(project, tmp_path):
    out, _ = build(project, tmp_path, page_size="Letter", margin_left=2.54)
    section = docx.Document(str(out)).sections[0]
    assert abs(section.page_width.cm - 21.59) < 0.01
    assert abs(section.left_margin.cm - 2.54) < 0.01
    assert "PAGE" in xml(out, "word/footer1.xml")


def test_estimate_pages_grows_with_content(make_project, tmp_path):
    small = make_project({"a.py": "x = 1\n" * 10}, name="small")
    big = make_project({"a.py": "x = 1\n" * 2000}, name="big")
    cfg = Config()
    est = []
    for root in (small, big):
        result = scan(root, cfg)
        files = load_all(result, cfg)
        est.append(estimate_pages(cfg, files, None))
    assert est[0] == 1
    assert 30 <= est[1] <= 50  # ~2000 lines at ~55 lines per A4 page
